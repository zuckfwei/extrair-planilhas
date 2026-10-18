"""Spreadsheet readers and streaming TXT export, independent from the desktop UI."""
from dataclasses import dataclass
from pathlib import Path
from time import perf_counter
import os
import tempfile

import openpyxl
from openpyxl.utils import get_column_letter
import xlrd
from odf import teletype
from odf.opendocument import load
from odf.table import Table, TableRow
from odf.text import P

SUPPORTED = {'.xlsx', '.xlsm', '.xls', '.ods'}
PREVIEW_LIMIT = 1000


class Cancelled(Exception):
    pass


def check_cancel(cancel):
    if cancel and cancel.is_set():
        raise Cancelled('Operação cancelada.')


class Spreadsheet:
    """One open workbook, reused when changing sheets; never executes macros."""
    def __init__(self, path):
        self.path = Path(path)
        self.ext = self.path.suffix.lower()
        if self.ext not in SUPPORTED:
            raise ValueError('Escolha um arquivo XLSX, XLSM, XLS ou ODS.')
        if self.ext in {'.xlsx', '.xlsm'}:
            self.book = openpyxl.load_workbook(self.path, read_only=True, data_only=False)
            self.sheets = self.book.sheetnames
        elif self.ext == '.xls':
            self.book = xlrd.open_workbook(str(self.path), on_demand=True)
            self.sheets = self.book.sheet_names()
        else:
            self.book = load(str(self.path))
            self.tables = {s.getAttribute('name'): s for s in self.book.spreadsheet.getElementsByType(Table)}
            self.sheets = list(self.tables)
        if not self.sheets:
            self.close()
            raise ValueError('Esta planilha não contém abas.')

    def close(self):
        if self.ext in {'.xlsx', '.xlsm'}:
            self.book.close()
        elif self.ext == '.xls':
            self.book.release_resources()

    def cells(self, sheet, formulas=True, cancel=None):
        if sheet not in self.sheets:
            raise ValueError('A aba selecionada não existe.')
        if self.ext in {'.xlsx', '.xlsm'}:
            book = self.book if formulas else openpyxl.load_workbook(self.path, read_only=True, data_only=True)
            try:
                for row in book[sheet].iter_rows():
                    check_cancel(cancel)
                    for cell in row:
                        if cell.value is not None:
                            yield cell.coordinate, str(cell.value), cell.data_type == 'f'
            finally:
                if book is not self.book:
                    book.close()
        elif self.ext == '.xls':
            ws = self.book.sheet_by_name(sheet)
            for r in range(ws.nrows):
                check_cancel(cancel)
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    if cell.ctype not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        yield f'{get_column_letter(c + 1)}{r + 1}', str(cell.value), False
        else:
            row_idx = 1
            for row in self.tables[sheet].getElementsByType(TableRow):
                check_cancel(cancel)
                repeat = int(row.getAttribute('numberrowsrepeated') or 1)
                entries = []
                col_idx = 1
                for cell in row.childNodes:
                    tag = getattr(cell, 'tagName', '')
                    if tag not in {'table:table-cell', 'table:covered-table-cell'}:
                        continue
                    cols = int(cell.getAttribute('numbercolumnsrepeated') or 1)
                    if tag == 'table:table-cell':
                        formula = cell.getAttribute('formula')
                        text = '\n'.join(teletype.extractText(p) for p in cell.getElementsByType(P))
                        if not text:
                            text = next((cell.getAttribute(a) for a in ('value', 'datevalue', 'timevalue', 'booleanvalue', 'stringvalue') if cell.getAttribute(a) is not None), '')
                        value = formula if formulas and formula else text
                        if value != '':
                            entries.append((col_idx, cols, value, bool(formulas and formula)))
                    col_idx += cols
                # Skip repeated empty rows/columns without expanding them.
                if entries:
                    for r in range(row_idx, row_idx + repeat):
                        check_cancel(cancel)
                        for start, cols, value, formula in entries:
                            for c in range(start, start + cols):
                                if c % 1024 == 0:
                                    check_cancel(cancel)
                                yield f'{get_column_letter(c)}{r}', str(value), formula
                row_idx += repeat


@dataclass
class Extraction:
    path: Path
    preview: list
    count: int
    formulas: int
    seconds: float
    size: int

    def cleanup(self):
        self.path.unlink(missing_ok=True)


def extract(workbook, sheet, formulas=True, cancel=None, progress=None):
    """Keep only a limited preview in memory; write the complete result to disk."""
    start = perf_counter()
    fd, name = tempfile.mkstemp(prefix='extrair-planilhas-', suffix='.txt')
    path = Path(name)
    preview = []
    count = formula_count = 0
    try:
        with os.fdopen(fd, 'w', encoding='utf-8', newline='\n') as output:
            for coord, value, formula in workbook.cells(sheet, formulas, cancel):
                check_cancel(cancel)
                if count:
                    output.write('\n')
                output.write(f'{coord}:{value}')
                if len(preview) < PREVIEW_LIMIT:
                    preview.append((coord, value))
                count += 1
                formula_count += formula
                if progress and count % 2000 == 0:
                    progress(count)
        check_cancel(cancel)
        return Extraction(path, preview, count, formula_count, perf_counter() - start, path.stat().st_size)
    except BaseException:
        path.unlink(missing_ok=True)
        raise


def export_text(source, destination, cancel=None):
    """Atomic export: preserve an existing destination if copying fails/cancels."""
    destination = Path(destination)
    fd, temporary = tempfile.mkstemp(prefix='.extrair-', suffix='.tmp', dir=destination.parent)
    try:
        with os.fdopen(fd, 'wb') as output, Path(source).open('rb') as input_file:
            while True:
                check_cancel(cancel)
                chunk = input_file.read(1024 * 1024)
                if not chunk:
                    break
                output.write(chunk)
        check_cancel(cancel)
        os.replace(temporary, destination)
        return destination
    finally:
        Path(temporary).unlink(missing_ok=True)
