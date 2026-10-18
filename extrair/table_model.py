"""Modelo de grade usado na prévia da planilha."""
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from PySide6.QtCore import QAbstractTableModel, QModelIndex, Qt
from PySide6.QtGui import QColor


class TableModel(QAbstractTableModel):
    def __init__(self):
        super().__init__()
        self.cells = {}
        self.columns = []
        self.rows = []
        self.visible_rows = []

    def set_cells(self, cells):
        self.beginResetModel()
        self.cells = {}
        used_columns = set()
        used_rows = set()

        for coordinate, value in cells:
            column_name, row = coordinate_from_string(coordinate)
            column = column_index_from_string(column_name)
            self.cells[row, column] = value
            used_columns.add(column)
            used_rows.add(row)

        if used_columns and max(used_columns) - min(used_columns) <= 100:
            self.columns = list(range(min(used_columns), max(used_columns) + 1))
        else:
            self.columns = sorted(used_columns)
        self.rows = sorted(used_rows)
        self.visible_rows = self.rows.copy()
        self.endResetModel()

    def set_filter(self, text):
        query = text.casefold().strip()
        self.beginResetModel()
        if not query:
            self.visible_rows = self.rows.copy()
        else:
            self.visible_rows = [
                row for row in self.rows
                if any(
                    query in f'{get_column_letter(column)}{row}'.casefold()
                    or query in str(self.cells.get((row, column), '')).casefold()
                    for column in self.columns
                )
            ]
        self.endResetModel()

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.visible_rows)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.columns)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        row = self.visible_rows[index.row()]
        column = self.columns[index.column()]
        value = self.cells.get((row, column), '')
        if role in (Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.ToolTipRole):
            return value
        if role == Qt.ItemDataRole.ForegroundRole and value:
            return QColor('#263b34')
        if role == Qt.ItemDataRole.TextAlignmentRole:
            return Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role != Qt.ItemDataRole.DisplayRole:
            return None
        if orientation == Qt.Orientation.Horizontal and section < len(self.columns):
            return get_column_letter(self.columns[section])
        if orientation == Qt.Orientation.Vertical and section < len(self.visible_rows):
            return str(self.visible_rows[section])
        return None
