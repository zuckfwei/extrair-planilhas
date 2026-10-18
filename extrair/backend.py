"""Leitura de planilhas e exportação de texto, sem dependências da interface."""
from collections.abc import Callable, Iterator
from dataclasses import dataclass
import os
from pathlib import Path
import tempfile
from threading import Event
from time import perf_counter

import openpyxl
from openpyxl.utils import get_column_letter
import xlrd
from odf import teletype
from odf.opendocument import load
from odf.table import Table, TableRow
from odf.text import P


SUPPORTED_FORMATS = {'.xlsx', '.xlsm', '.xls', '.ods'}
PREVIEW_LIMIT = 1000
COPY_CHUNK_SIZE = 1024 * 1024
ODS_VALUE_ATTRIBUTES = ('value', 'datevalue', 'timevalue', 'booleanvalue', 'stringvalue')

Cell = tuple[str, str, bool]
ProgressCallback = Callable[[int], None]


class Cancelled(Exception):
    """Indica que o usuário cancelou uma operação em andamento."""


def ensure_not_cancelled(cancel_event: Event | None) -> None:
    if cancel_event is None:
        return
    if cancel_event.is_set():
        raise Cancelled('Operação cancelada.')


class Spreadsheet:
    """Mantém uma planilha aberta enquanto o usuário alterna entre abas."""

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.extension = self.path.suffix.lower()
        self.workbook = None
        self.ods_tables = {}
        self.sheet_names = []

        if self.extension not in SUPPORTED_FORMATS:
            raise ValueError('Escolha um arquivo XLSX, XLSM, XLS ou ODS.')

        self._open()
        if self.sheet_names:
            return

        self.close()
        raise ValueError('Esta planilha não contém abas.')

    def _open(self) -> None:
        if self.extension in {'.xlsx', '.xlsm'}:
            self.workbook = openpyxl.load_workbook(
                self.path,
                read_only=True,
                data_only=False,
            )
            self.sheet_names = list(self.workbook.sheetnames)
            return

        if self.extension == '.xls':
            self.workbook = xlrd.open_workbook(str(self.path), on_demand=True)
            self.sheet_names = self.workbook.sheet_names()
            return

        self.workbook = load(str(self.path))
        tables = self.workbook.spreadsheet.getElementsByType(Table)
        for table in tables:
            name = table.getAttribute('name')
            self.ods_tables[name] = table
        self.sheet_names = list(self.ods_tables)

    def close(self) -> None:
        if self.extension in {'.xlsx', '.xlsm'}:
            self.workbook.close()
            return
        if self.extension == '.xls':
            self.workbook.release_resources()

    def iter_cells(
        self,
        sheet_name: str,
        include_formulas: bool = True,
        cancel_event: Event | None = None,
    ) -> Iterator[Cell]:
        if sheet_name not in self.sheet_names:
            raise ValueError('A aba selecionada não existe.')

        if self.extension in {'.xlsx', '.xlsm'}:
            yield from self._iter_openxml_cells(sheet_name, include_formulas, cancel_event)
            return
        if self.extension == '.xls':
            yield from self._iter_xls_cells(sheet_name, cancel_event)
            return
        yield from self._iter_ods_cells(sheet_name, include_formulas, cancel_event)

    def _iter_openxml_cells(
        self,
        sheet_name: str,
        include_formulas: bool,
        cancel_event: Event | None,
    ) -> Iterator[Cell]:
        selected_workbook = self.workbook
        opened_cached_workbook = False

        if not include_formulas:
            selected_workbook = openpyxl.load_workbook(
                self.path,
                read_only=True,
                data_only=True,
            )
            opened_cached_workbook = True

        try:
            worksheet = selected_workbook[sheet_name]
            for row in worksheet.iter_rows():
                ensure_not_cancelled(cancel_event)
                for cell in row:
                    if cell.value is None:
                        continue
                    is_formula = cell.data_type == 'f'
                    yield cell.coordinate, str(cell.value), is_formula
        finally:
            if opened_cached_workbook:
                selected_workbook.close()

    def _iter_xls_cells(
        self,
        sheet_name: str,
        cancel_event: Event | None,
    ) -> Iterator[Cell]:
        worksheet = self.workbook.sheet_by_name(sheet_name)
        empty_types = {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}

        for row_index in range(worksheet.nrows):
            ensure_not_cancelled(cancel_event)
            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                if cell.ctype in empty_types:
                    continue
                column_name = get_column_letter(column_index + 1)
                coordinate = f'{column_name}{row_index + 1}'
                yield coordinate, str(cell.value), False

    def _iter_ods_cells(
        self,
        sheet_name: str,
        include_formulas: bool,
        cancel_event: Event | None,
    ) -> Iterator[Cell]:
        current_row = 1
        table = self.ods_tables[sheet_name]

        for row in table.getElementsByType(TableRow):
            ensure_not_cancelled(cancel_event)
            repeated_rows = _repeat_count(row, 'numberrowsrepeated')
            entries = _read_ods_row(row, include_formulas)

            if not entries:
                current_row += repeated_rows
                continue

            for row_offset in range(repeated_rows):
                row_number = current_row + row_offset
                yield from _expand_ods_entries(entries, row_number, cancel_event)

            current_row += repeated_rows


def _repeat_count(element, attribute: str) -> int:
    value = element.getAttribute(attribute)
    if not value:
        return 1
    return int(value)


def _read_ods_row(row, include_formulas: bool) -> list[tuple[int, int, str, bool]]:
    entries = []
    current_column = 1

    for cell in row.childNodes:
        tag_name = getattr(cell, 'tagName', '')
        if tag_name not in {'table:table-cell', 'table:covered-table-cell'}:
            continue

        repeated_columns = _repeat_count(cell, 'numbercolumnsrepeated')
        if tag_name == 'table:table-cell':
            value, is_formula = _read_ods_cell(cell, include_formulas)
            if value != '':
                entries.append((current_column, repeated_columns, value, is_formula))
        current_column += repeated_columns

    return entries


def _read_ods_cell(cell, include_formulas: bool) -> tuple[str, bool]:
    formula = cell.getAttribute('formula')
    if include_formulas and formula:
        return formula, True

    paragraphs = []
    for paragraph in cell.getElementsByType(P):
        paragraphs.append(teletype.extractText(paragraph))
    text = '\n'.join(paragraphs)
    if text:
        return text, False

    for attribute in ODS_VALUE_ATTRIBUTES:
        value = cell.getAttribute(attribute)
        if value is not None:
            return value, False
    return '', False


def _expand_ods_entries(
    entries: list[tuple[int, int, str, bool]],
    row_number: int,
    cancel_event: Event | None,
) -> Iterator[Cell]:
    for start_column, repeated_columns, value, is_formula in entries:
        end_column = start_column + repeated_columns
        for column_number in range(start_column, end_column):
            if column_number % 1024 == 0:
                ensure_not_cancelled(cancel_event)
            column_name = get_column_letter(column_number)
            yield f'{column_name}{row_number}', str(value), is_formula


@dataclass
class ExtractionResult:
    temporary_path: Path
    preview_cells: list[tuple[str, str]]
    cell_count: int
    formula_count: int
    elapsed_seconds: float
    file_size: int

    def remove_temporary_file(self) -> None:
        self.temporary_path.unlink(missing_ok=True)


def extract_sheet(
    spreadsheet: Spreadsheet,
    sheet_name: str,
    include_formulas: bool = True,
    cancel_event: Event | None = None,
    progress_callback: ProgressCallback | None = None,
) -> ExtractionResult:
    """Grava o resultado completo em disco e mantém uma prévia na memória."""
    started_at = perf_counter()
    descriptor, temporary_name = tempfile.mkstemp(
        prefix='extrair-planilhas-',
        suffix='.txt',
    )
    temporary_path = Path(temporary_name)
    preview_cells = []
    cell_count = 0
    formula_count = 0

    try:
        with os.fdopen(descriptor, 'w', encoding='utf-8', newline='\n') as output:
            cells = spreadsheet.iter_cells(sheet_name, include_formulas, cancel_event)
            for coordinate, value, is_formula in cells:
                ensure_not_cancelled(cancel_event)
                if cell_count > 0:
                    output.write('\n')
                output.write(f'{coordinate}:{value}')

                if len(preview_cells) < PREVIEW_LIMIT:
                    preview_cells.append((coordinate, value))
                cell_count += 1
                if is_formula:
                    formula_count += 1

                should_report_progress = cell_count % 2000 == 0
                if progress_callback is not None and should_report_progress:
                    progress_callback(cell_count)

        ensure_not_cancelled(cancel_event)
        elapsed_seconds = perf_counter() - started_at
        file_size = temporary_path.stat().st_size
        return ExtractionResult(
            temporary_path=temporary_path,
            preview_cells=preview_cells,
            cell_count=cell_count,
            formula_count=formula_count,
            elapsed_seconds=elapsed_seconds,
            file_size=file_size,
        )
    except BaseException:
        temporary_path.unlink(missing_ok=True)
        raise


def export_text_file(
    source: str | Path,
    destination: str | Path,
    cancel_event: Event | None = None,
) -> Path:
    """Substitui o destino somente depois que a cópia termina com sucesso."""
    source_path = Path(source)
    destination_path = Path(destination)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix='.extrair-',
        suffix='.tmp',
        dir=destination_path.parent,
    )
    temporary_path = Path(temporary_name)

    try:
        with os.fdopen(descriptor, 'wb') as output:
            with source_path.open('rb') as input_file:
                while True:
                    ensure_not_cancelled(cancel_event)
                    chunk = input_file.read(COPY_CHUNK_SIZE)
                    if not chunk:
                        break
                    output.write(chunk)

        ensure_not_cancelled(cancel_event)
        os.replace(temporary_path, destination_path)
        return destination_path
    finally:
        temporary_path.unlink(missing_ok=True)
